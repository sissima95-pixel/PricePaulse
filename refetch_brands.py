"""Re-fetch brand for ASINs whose brand is empty or a descriptive word.
Writes an updated Excel next to the input. Usage:
    python refetch_brands.py test_output_us/asin_detail_20260811_111615.xlsx US
"""
import sys, time
import openpyxl
from pricepulse.fetcher import MarketSession

DESC = {"portable","wireless","mini","smart","solar","outdoor","indoor","electric",
        "rechargeable","handheld","cordless","neck","desk","wall","car","baby","home",
        "travel","set","kit","usb","led","personal","small","large","dual","unknown"}

src = sys.argv[1]
market = sys.argv[2] if len(sys.argv) > 2 else "US"
wb = openpyxl.load_workbook(src)
ws = wb.active
hdr = [str(c.value) for c in ws[1]]
ci = {h: i + 1 for i, h in enumerate(hdr)}  # 1-based

targets = []
for row in range(2, ws.max_row + 1):
    status = ws.cell(row, ci['Status']).value
    brand = (ws.cell(row, ci['Brand']).value or "").strip()
    if status == 'ok' and (not brand or brand.lower() in DESC):
        targets.append(row)
print(f'{len(targets)} ASINs need brand refetch on {market}')

s = MarketSession(market)
s.prepare()
fixed = 0
t0 = time.time()
for n, row in enumerate(targets, 1):
    asin = ws.cell(row, ci['ASIN']).value
    try:
        r = s.fetch(asin)
        if r.brand:
            ws.cell(row, ci['Brand']).value = r.brand
            fixed += 1
        print(f'  [{n}/{len(targets)}] {asin} -> {r.brand or "(still unknown)"}')
    except Exception as e:
        print(f'  [{n}/{len(targets)}] {asin} -> ERROR {e}')
    time.sleep(0.6)

out = src.replace('.xlsx', '_brandfix.xlsx')
wb.save(out)
print(f'\nFixed {fixed}/{len(targets)} in {time.time()-t0:.0f}s -> {out}')

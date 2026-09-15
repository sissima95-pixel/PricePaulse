# -*- coding: utf-8 -*-
"""ASIN PricePulse reporter — generates Excel and interactive HTML reports.

v1.2.1 — 2026-08-11. Design system overhaul: Indigo/Violet palette, Inter font,
refined spacing, shadows, and component styles per design spec.
"""
from __future__ import annotations

import json
import html as html_mod
from pathlib import Path
from dataclasses import asdict
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from .fetcher import PriceResult


# ===========================================================================
# Excel output
# ===========================================================================

def write_csv(results: list, path) -> None:
    """Write results to styled Excel file (.xlsx)."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        path = path.with_suffix(".xlsx")
    if not results:
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["No results"])
        wb.save(path)
        return

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    fieldnames = [
        "asin", "market", "currency", "price", "display_price",
        "title", "brand", "availability", "status", "url", "error",
        "search_rank", "purchase_rank", "search_volume",
    ]
    header_names = [
        "ASIN", "Market", "Currency", "Price", "Display Price",
        "Title", "Brand", "Availability", "Status", "URL", "Error",
        "Search Rank", "Purchase Rank", "Search Volume",
    ]
    series_keys: list[str] = []
    for r in results:
        d = asdict(r) if not isinstance(r, dict) else r
        vs = d.get("volume_series")
        if vs and isinstance(vs, dict):
            for k in vs:
                if k not in series_keys:
                    series_keys.append(k)
    series_keys.sort()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ASIN Detail"
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_font = Font(name="Segoe UI", size=10)
    body_align = Alignment(vertical="center", wrap_text=False)
    thin_border = Border(
        left=Side(style="thin", color="E5E5E5"),
        right=Side(style="thin", color="E5E5E5"),
        top=Side(style="thin", color="E5E5E5"),
        bottom=Side(style="thin", color="E5E5E5"),
    )
    headers = header_names + series_keys
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    for r in results:
        d = asdict(r) if not isinstance(r, dict) else r
        row = [d.get(fn, "") for fn in fieldnames]
        vs = d.get("volume_series") or {}
        for sk in series_keys:
            row.append(vs.get(sk, ""))
        ws.append(row)
    alt_fill = PatternFill(start_color="FAFAFA", end_color="FAFAFA", fill_type="solid")
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = body_font
            cell.alignment = body_align
            cell.border = thin_border
        if row_idx % 2 == 0:
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row=row_idx, column=col_idx).fill = alt_fill
    for col_idx in range(1, ws.max_column + 1):
        max_len = len(str(headers[col_idx - 1]))
        for row_idx in range(2, min(ws.max_row + 1, 20)):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, min(len(str(val)), 40))
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max_len + 3
    ws.freeze_panes = "A2"
    wb.save(path)


# ===========================================================================
# HTML output
# ===========================================================================

_CHART_COLORS = ['#6366f1', '#8b5cf6', '#ec4899', '#f59e0b', '#10b981',
                 '#06b6d4', '#f97316', '#64748b', '#3b82f6', '#22c55e']
_TIER_COLORS = ['#10b981', '#3b82f6', '#ec4899']  # green, blue, pink
_TIER_BG = ['#dcfce7', '#dbeafe', '#fce7f3']
_TIER_TEXT = ['#166534', '#1d4ed8', '#9f1239']


def write_html(results: list, path, title: str = "ASIN Price Intelligence Report",
               subtitle: str = "") -> None:
    path = Path(path)
    if not results:
        path.write_text("<!DOCTYPE html><html><body><p>No results</p></body></html>", encoding="utf-8")
        return
    rows = [asdict(r) if not isinstance(r, dict) else r for r in results]

    # Post-process: clear brands that are actually descriptive words (not real brands)
    _desc_words = {
        "portable", "wireless", "mini", "smart", "solar", "outdoor",
        "indoor", "electric", "digital", "automatic", "universal",
        "rechargeable", "bluetooth", "waterproof", "adjustable",
        "foldable", "handheld", "cordless", "stainless", "heavy",
        "professional", "personal", "upgraded", "advanced", "ultra",
        "super", "extra", "large", "small", "big", "slim", "compact",
        "powerful", "quiet", "silent", "fast", "quick", "high", "low",
        "dual", "double", "single", "multi",
        "neck", "desk", "wall", "floor", "table", "car", "bed",
        "baby", "pet", "dog", "cat", "kids", "home", "kitchen",
        "garden", "office", "camping", "travel", "sport", "gym",
        "set", "kit", "pair", "usb", "led", "lcd", "hd", "wifi",
    }
    for r in rows:
        b = (r.get("brand") or "").strip()
        if b.lower() in _desc_words:
            r["brand"] = ""
    has_sr = any(r.get("search_rank") is not None for r in rows)
    has_pr = any(r.get("purchase_rank") is not None for r in rows)
    has_sv = any(r.get("search_volume") is not None for r in rows)
    has_vs = any(r.get("volume_series") for r in rows)

    markets_data: dict[str, list[dict]] = {}
    for r in rows:
        markets_data.setdefault(r.get("market", "?"), []).append(r)

    panels_html = ""
    tabs_html = ""
    tier_bounds_map: dict[str, list] = {}  # market -> [[lo,hi], [lo,hi], [lo,hi]]
    first = True
    for mk, mrows in markets_data.items():
        active = " active" if first else ""
        tabs_html += f'<div class="tab{active}" data-market="{mk}">{mk} ({len(mrows)})</div>\n'
        panel_html, tier_bounds = _render_panel(mk, mrows, active, has_sr, has_pr, has_sv)
        panels_html += panel_html
        tier_bounds_map[mk] = tier_bounds
        first = False

    json_rows = [{k: v for k, v in r.items() if k != "volume_series"} for r in rows]
    vol_json = {}
    if has_vs:
        for r in rows:
            vs = r.get("volume_series")
            if vs:
                vol_json[r["asin"]] = vs

    html = _full_html(html_mod.escape(title), html_mod.escape(subtitle) if subtitle else "",
                      tabs_html, panels_html, json_rows, vol_json, tier_bounds_map,
                      has_sr, has_pr, has_sv, has_vs)
    path.write_text(html, encoding="utf-8")


def _price_bands(prices: list[float]) -> list[tuple[str, int, float, float]]:
    """Smart 3-tier price banding using tercile (33rd/67th percentile) splits.

    This ensures each tier has a roughly balanced number of ASINs,
    making the segmentation meaningful regardless of distribution skew.
    Boundaries are rounded to "nice" numbers for readability.
    Returns [(label, count, lo, hi), ...]
    """
    if not prices:
        return []
    sorted_p = sorted(prices)
    n = len(sorted_p)
    mn, mx = sorted_p[0], sorted_p[-1]

    if mx - mn < 1:
        # All same price — single tier
        return [("Entry", n, mn, mx)]

    # Tercile boundaries (33rd and 67th percentile)
    b1_raw = sorted_p[n // 3]
    b2_raw = sorted_p[2 * n // 3]

    # Round to "nice" numbers for readability
    def _nice(v):
        if v <= 10:
            return round(v)
        elif v <= 100:
            return round(v / 5) * 5       # round to nearest 5
        elif v <= 500:
            return round(v / 10) * 10      # round to nearest 10
        else:
            return round(v / 50) * 50      # round to nearest 50

    b1 = _nice(b1_raw)
    b2 = _nice(b2_raw)
    # Ensure boundaries are distinct and ordered
    if b1 <= mn:
        b1 = _nice(mn + (mx - mn) * 0.33)
    if b2 <= b1:
        b2 = _nice(b1 + (mx - b1) * 0.5)
    if b2 >= mx:
        b2 = _nice(mn + (mx - mn) * 0.67)
    if b1 >= b2:
        # Fallback: equal-width
        step = (mx - mn) / 3
        b1 = _nice(mn + step)
        b2 = _nice(mn + step * 2)

    labels = ["Entry", "Mid-tier", "Premium"]
    bounds = [(0, b1), (b1, b2), (b2, float('inf'))]
    bands = []
    for i, (lo, hi) in enumerate(bounds):
        if i == 2:
            count = sum(1 for p in prices if p >= lo)
        else:
            count = sum(1 for p in prices if lo <= p < hi)
        bands.append((labels[i], count, lo if i == 0 else lo, hi if i < 2 else mx))
    return bands


def _render_panel(mk, mrows, active, has_sr, has_pr, has_sv):
    """Returns (html_string, tier_bounds_list)."""
    ok_rows = [r for r in mrows if r.get("status") == "ok" and r.get("price")]
    prices = [r["price"] for r in ok_rows]
    bands = _price_bands(prices)
    cur = ok_rows[0]["currency"] if ok_rows else ""
    total_priced = sum(c for _, c, _, _ in bands)
    display = "block" if active else "none"

    # Collect tier bounds for JS sync
    tier_bounds = []
    for label, count, lo, hi in bands:
        tier_bounds.append({"name": label, "lo": lo, "hi": hi})

    # --- Helper: which tier does a price belong to ---
    def _get_tier(p):
        for i, (lbl, _, lo, hi) in enumerate(bands):
            if i == len(bands) - 1:
                if p >= lo:
                    return lbl
            else:
                if lo <= p < hi:
                    return lbl
        return bands[-1][0] if bands else ""

    # --- Tier cards ---
    tier_cards = ""
    for i, (label, count, lo, hi) in enumerate(bands):
        pct = f"{count/total_priced*100:.0f}" if total_priced else "0"
        tier_prices = [r["price"] for r in ok_rows if _get_tier(r["price"]) == label]
        avg_p = sum(tier_prices) / len(tier_prices) if tier_prices else 0
        range_str = f"${lo:.0f}\u2013${hi:.0f}" if i < 2 else f"${lo:.0f}+"
        tier_cards += f'''
        <div class="tier-card clickable-bar" data-filter-type="price" data-filter-value="{label}" data-market="{mk}">
          <div class="tier-badge" style="background:{_TIER_BG[i]};color:{_TIER_TEXT[i]}">{label}</div>
          <div class="tier-range">{range_str}</div>
          <div class="tier-num">{count} <span class="tier-pct">({pct}%)</span></div>
          <div class="tier-avg">Avg ${avg_p:.2f}</div>
        </div>'''

    # --- Brand data ---
    brand_counts: dict[str, int] = {}
    brand_prices: dict[str, list[float]] = {}
    brand_sv: dict[str, float] = {}  # total search volume
    brand_pr: dict[str, float] = {}  # total purchase rank (lower = better, we use count of ASINs)
    for r in ok_rows:
        b = r.get("brand", "").strip() or "Unknown"
        brand_counts[b] = brand_counts.get(b, 0) + 1
        brand_prices.setdefault(b, []).append(r["price"])
        sv = r.get("search_volume") or 0
        brand_sv[b] = brand_sv.get(b, 0) + sv
        # Purchase rank: count ASINs with purchase_rank (as proxy for purchase share)
        pr = r.get("purchase_rank")
        if pr is not None:
            brand_pr[b] = brand_pr.get(b, 0) + 1

    # Brand avg price chart (Top 12)
    brand_avg = sorted([(b, sum(ps)/len(ps), len(ps)) for b, ps in brand_prices.items()],
                       key=lambda x: -x[2])[:12]
    max_avg = max((a for _, a, _ in brand_avg), default=1)
    brand_avg_html = '<div class="card"><div class="card-title">Brand Avg Price \u00b7 Top 12</div>'
    for i, (bname, avg, cnt) in enumerate(brand_avg):
        pct = avg / max_avg * 100
        color = _CHART_COLORS[i % len(_CHART_COLORS)]
        brand_avg_html += f'''<div class="hbar-row clickable-bar" data-filter-type="brand" data-filter-value="{html_mod.escape(bname)}" data-market="{mk}">
          <div class="hbar-label">{html_mod.escape(bname)}</div>
          <div class="hbar-track"><div class="hbar-fill" style="width:{pct:.1f}%;background:{color}"></div></div>
          <div class="hbar-val">${avg:.0f} <span class="hbar-cnt">({cnt})</span></div>
        </div>'''
    brand_avg_html += '</div>'

    # --- Brand Share donut charts (JS-rendered) ---
    # We pass brand_sv and brand_pr as JSON for JS to render donuts
    # Placeholder divs
    donut_html = ""
    if has_sv:
        donut_html += f'<div class="card"><div class="card-title">BRAND SHARE BY SEARCH VOLUME</div><canvas id="donut-sv-{mk}" height="260"></canvas></div>'
    donut_html += f'<div class="card"><div class="card-title">BRAND SHARE BY ASIN COUNT</div><canvas id="donut-cnt-{mk}" height="260"></canvas></div>'

    # --- Brand Share by Price Tier (collapsible) ---
    tier_brand_html = '<div class="card">'
    for i, (label, count, lo, hi) in enumerate(bands):
        range_str = f"${lo:.0f}\u2013${hi:.0f}" if i < 2 else f"${lo:.0f}+"
        tier_brand_html += f'<div class="tier-brand-section">'
        tier_brand_html += f'<div class="tier-brand-header" style="border-left:4px solid {_TIER_COLORS[i]}">'
        tier_brand_html += f'<span class="tier-badge" style="background:{_TIER_BG[i]};color:{_TIER_TEXT[i]}">{label}</span> {range_str}'
        tier_brand_html += f'</div>'
        # Brands in this tier
        tier_rows = [r for r in ok_rows if _get_tier(r["price"]) == label]
        tier_brands: dict[str, list] = {}
        for r in tier_rows:
            b = r.get("brand", "").strip() or "Unknown"
            tier_brands.setdefault(b, []).append(r)
        sorted_tb = sorted(tier_brands.items(), key=lambda x: -len(x[1]))
        for bname, basins in sorted_tb:
            asin_list = ", ".join(r.get("asin", "") for r in basins[:10])
            extra = f" +{len(basins)-10} more" if len(basins) > 10 else ""
            tier_brand_html += f'<details class="tier-brand-detail">'
            tier_brand_html += f'<summary><span class="tb-brand">{html_mod.escape(bname)}</span> <span class="tb-count">{len(basins)}</span></summary>'
            tier_brand_html += f'<div class="tb-asins">{html_mod.escape(asin_list)}{extra}</div>'
            tier_brand_html += f'</details>'
        tier_brand_html += '</div>'
    tier_brand_html += '</div>'

    # --- INSIGHTS (professional format: EN + CN italic) ---
    # A. PRICE OVERVIEW
    tier_insight = ""
    if bands and total_priced:
        sorted_bands = sorted(bands, key=lambda x: -x[1])
        dominant = sorted_bands[0]
        dom_pct = dominant[1] / total_priced * 100
        smallest = sorted_bands[-1]
        sm_pct = smallest[1] / total_priced * 100
        avg_price = sum(prices) / len(prices)
        median_price = sorted(prices)[len(prices) // 2]

        tier_insight = '<div class="insight-box">'
        tier_insight += '<div class="insight-title">A. PRICE OVERVIEW / \u4ef7\u683c\u6982\u89c8</div>'
        # EN
        tier_insight += f'<p>Price range spans <b>${min(prices):.0f}\u2013${max(prices):.0f}</b> ({cur}) '
        tier_insight += f'with median <b>${median_price:.0f}</b> and mean <b>${avg_price:.0f}</b>. '
        tier_insight += f'The <b>{dominant[0]}</b> tier dominates at <b>{dom_pct:.0f}%</b> of ASINs ({dominant[1]}), '
        tier_insight += f'while <b>{smallest[0]}</b> holds <b>{sm_pct:.0f}%</b> ({smallest[1]}).'
        if dom_pct > 50:
            tier_insight += f' Heavy concentration in the {dominant[0]} segment suggests '
            tier_insight += f'potential whitespace in the {smallest[0]} tier.'
        tier_insight += '</p>'
        # CN
        tier_insight += f'<p class="cn">\u4ef7\u683c\u8303\u56f4${min(prices):.0f}\u2013${max(prices):.0f}({cur})\uff0c'
        tier_insight += f'\u4e2d\u4f4d\u6570${median_price:.0f}\uff0c\u5747\u4ef7${avg_price:.0f}\u3002'
        tier_insight += f'{dominant[0]}\u6863\u4f4d\u5360\u6bd4{dom_pct:.0f}%({dominant[1]}\u4e2aASIN)\uff0c'
        tier_insight += f'{smallest[0]}\u6863\u4f4d\u4ec5{sm_pct:.0f}%({smallest[1]}\u4e2a)\u3002'
        if dom_pct > 50:
            tier_insight += f'\u5e02\u573a\u91cd\u5ea6\u96c6\u4e2d\u5728{dominant[0]}\u6bb5\uff0c'
            tier_insight += f'{smallest[0]}\u6bb5\u5b58\u5728\u5dee\u5f02\u5316\u7a7a\u95f4\u3002'
        tier_insight += '</p>'
        # Banding note
        tier_insight += '<div class="insight-note">Price tiers split by tercile (33rd/67th percentile) '
        tier_insight += 'for balanced distribution. Boundaries rounded for readability.</div>'
        tier_insight += '</div>'

    # B. BRAND LANDSCAPE
    dist_insight = ""
    if brand_avg:
        n_brands = len(brand_counts)
        top3 = brand_avg[:3]
        top3_names = " + ".join(html_mod.escape(b) for b, _, _ in top3)
        top3_share = sum(c for _, _, c in top3) / total_priced * 100 if total_priced else 0
        top_brand = brand_avg[0]
        cheapest_brand = min(brand_avg, key=lambda x: x[1])

        dist_insight = '<div class="insight-box">'
        dist_insight += '<div class="insight-title">B. BRAND LANDSCAPE / \u54c1\u724c\u683c\u5c40</div>'
        # EN
        dist_insight += f'<p><b>{n_brands}</b> active brands. '
        dist_insight += f'Top 3 ({top3_names}) account for <b>{top3_share:.0f}%</b> of ASINs. '
        dist_insight += f'<b>{html_mod.escape(top_brand[0])}</b> leads with <b>{top_brand[2]}</b> ASINs '
        dist_insight += f'at avg <b>${top_brand[1]:.0f}</b>. '
        dist_insight += f'Most affordable brand: <b>{html_mod.escape(cheapest_brand[0])}</b> at avg <b>${cheapest_brand[1]:.0f}</b>.'
        if top3_share > 40:
            dist_insight += ' Moderate-to-high brand concentration.'
        elif top3_share < 20:
            dist_insight += ' Highly fragmented market with no dominant player.'
        dist_insight += '</p>'
        # CN
        dist_insight += f'<p class="cn">{n_brands}\u4e2a\u6d3b\u8dc3\u54c1\u724c\u3002'
        dist_insight += f'\u524d\u4e09({top3_names})\u5408\u8ba1\u5360{top3_share:.0f}%\u3002'
        dist_insight += f'{html_mod.escape(top_brand[0])}\u4ee5{top_brand[2]}\u4e2aASIN\u9886\u5148\uff0c'
        dist_insight += f'\u5747\u4ef7${top_brand[1]:.0f}\u3002'
        dist_insight += f'\u6700\u4f4e\u4ef7\u54c1\u724c:{html_mod.escape(cheapest_brand[0])}\uff0c\u5747\u4ef7${cheapest_brand[1]:.0f}\u3002'
        dist_insight += '</p></div>'

    # --- Volume chart placeholder ---
    vol_html = ""
    if has_sv:
        vol_html = f'''<div class="section">
          <div class="section-title">5. Search Volume by ASIN (Top 15) <span class="vol-filter-label" id="vol-filter-{mk}"></span></div>
          <div class="table-actions"><button class="action-btn" onclick="copyVolAsins('{mk}')" title="Copy ASINs">Copy ASINs</button></div>
          <div class="card"><div id="vol-chart-{mk}"></div></div>
        </div>'''

    # --- Table ---
    table_html = _render_table(mk, mrows, has_sr, has_pr)

    panel_str = f'''
    <div class="tab-content{" active" if active else ""}" data-market="{mk}" style="display:{display}">
      <div class="section">
        <div class="section-title">1. Price Tier Overview</div>
        <div class="tier-row">{tier_cards}</div>
        {tier_insight}
      </div>
      <div class="section">
        <div class="section-title">2. Distribution</div>
        <div class="grid-2">
          <div class="card tier-chart-card"><div class="card-title">ASIN COUNT BY PRICE TIER</div><div id="tier-svg-{mk}" class="tier-svg-wrap"></div></div>
          {brand_avg_html}
        </div>
      </div>
      <div class="section">
        <div class="section-title">3. Brand Share</div>
        <div class="grid-2">
          {donut_html}
        </div>
        <div class="card"><div class="bubble-cloud" id="bubble-{mk}"></div></div>
        {dist_insight}
      </div>
      <div class="section">
        <div class="section-title">4. Brand Share by Price Tier</div>
        {tier_brand_html}
      </div>
      {vol_html}
      <div class="section">
        <div class="section-title">6. ASIN Detail</div>
        <div class="table-actions">
          <button class="action-btn" onclick="copyAsins('{mk}')" title="Copy ASINs">Copy ASINs</button>
          <button class="action-btn" onclick="exportExcel('{mk}')" title="Export Excel">Export Excel</button>
        </div>
        {table_html}
      </div>
    </div>'''

    return (panel_str, tier_bounds)


def _render_table(mk, mrows, has_sr, has_pr):
    cols = ["ASIN", "Brand", "Title", "Price", "Status"]
    if has_sr:
        cols.append("Search Rank")
    if has_pr:
        cols.append("Purchase Rank")
    th = "".join(f'<th data-col="{i}">{c} <span class="th-sort" data-col="{i}">&#8597;</span> <span class="th-filter" data-col="{i}">\u25BC</span></th>' for i, c in enumerate(cols))
    tbody = ""
    for r in mrows:
        p = r.get("price")
        price_str = f'${p:.2f}' if p else "\u2014"
        status = r.get("status", "unknown")
        title_short = (r.get("title", "") or "")[:70]
        cells = [
            f'<a class="asin-link" href="{html_mod.escape(r.get("url",""))}" target="_blank">{html_mod.escape(r.get("asin",""))}</a>',
            html_mod.escape(r.get("brand", "") or ""),
            f'<span class="cell-title">{html_mod.escape(title_short)}</span>',
            f'<span class="price-tag">{price_str}</span>',
            f'<span class="st-{status}">{status}</span>',
        ]
        if has_sr:
            sv = r.get("search_rank")
            cells.append(f'<span class="rank-badge">{sv:.0f}</span>' if sv is not None else "\u2014")
        if has_pr:
            pv = r.get("purchase_rank")
            cells.append(f'<span class="rank-badge">{pv:.0f}</span>' if pv is not None else "\u2014")
        tbody += "<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>\n"

    return f'''<div class="table-wrap"><table id="table-{mk}"><thead><tr>{th}</tr></thead><tbody>{tbody}</tbody></table></div>'''


def _full_html(title, subtitle, tabs_html, panels_html, json_rows, vol_json, tier_bounds_map, has_sr, has_pr, has_sv, has_vs):
    row_json = json.dumps(json_rows, ensure_ascii=False, separators=(",", ":"))
    vol_j = json.dumps(vol_json, ensure_ascii=False, separators=(",", ":"))
    tier_j = json.dumps(tier_bounds_map, ensure_ascii=False, separators=(",", ":"))

    return f'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
<style>
:root{{
  --primary-500:#6366f1;--primary-600:#4f46e5;--primary-700:#4338ca;--primary-100:#e0e7ff;--primary-50:#eef2ff;
  --neutral-50:#fafafa;--neutral-100:#f5f5f5;--neutral-200:#e5e5e5;--neutral-400:#a3a3a3;--neutral-500:#737373;--neutral-600:#525252;--neutral-700:#404040;--neutral-800:#262626;--neutral-900:#171717;
  --success:#10b981;--warning:#f59e0b;--danger:#ef4444;
  --radius-md:8px;--radius-lg:12px;--radius-xl:14px;--radius-2xl:20px;
}}
*{{margin:0;padding:0;box-sizing:border-box;}}
body{{font-family:'Inter',-apple-system,BlinkMacSystemFont,'Segoe UI',system-ui,sans-serif;background:var(--neutral-50);color:var(--neutral-700);font-size:13px;line-height:1.6;}}
a{{color:var(--primary-600);text-decoration:none;}}
a:hover{{text-decoration:underline;}}

/* Hero */
.hero{{background:linear-gradient(135deg,#4f46e5 0%,#7c3aed 50%,#ec4899 100%);border-radius:var(--radius-2xl);padding:40px 44px;color:#fff;margin-bottom:28px;box-shadow:0 20px 40px rgba(79,70,229,0.25);position:relative;overflow:hidden;text-align:center;}}
.hero::before{{content:'';position:absolute;top:-50%;right:-20%;width:600px;height:600px;background:radial-gradient(circle,rgba(255,255,255,0.15) 0%,transparent 70%);pointer-events:none;}}
.hero h1{{font-size:26px;font-weight:700;letter-spacing:-0.02em;position:relative;}}
.hero .sub{{font-size:13px;opacity:.85;margin-top:6px;position:relative;}}

.container{{max-width:1400px;margin:0 auto;padding:32px 24px 100px;}}
.section{{margin-bottom:40px;}}
.section-title{{font-size:15px;font-weight:700;color:var(--neutral-800);border-left:4px solid var(--primary-500);padding-left:12px;margin-bottom:16px;}}

/* Tabs */
.tabs{{display:flex;gap:4px;margin-bottom:20px;border-bottom:2px solid var(--neutral-200);}}
.tab{{padding:8px 18px;font-size:12px;font-weight:600;cursor:pointer;border-radius:8px 8px 0 0;color:var(--neutral-500);border:1px solid transparent;border-bottom:none;position:relative;bottom:-2px;transition:all .2s ease;}}
.tab.active{{background:#fff;color:var(--primary-600);border-color:var(--neutral-200);border-bottom-color:#fff;}}
.tab:hover:not(.active){{background:var(--neutral-100);color:var(--neutral-800);}}
.tab-content{{display:none;}}
.tab-content.active{{display:block;}}

/* Tier cards */
.tier-row{{display:grid;grid-template-columns:repeat(auto-fit,minmax(240px,1fr));gap:16px;}}
.tier-card{{background:#fff;border-radius:var(--radius-xl);padding:20px 22px;box-shadow:0 2px 8px rgba(0,0,0,0.04);border:1px solid var(--neutral-200);cursor:pointer;transition:all .25s cubic-bezier(0.34,1.56,0.64,1);}}
.tier-card:hover{{transform:translateY(-3px);box-shadow:0 12px 24px rgba(0,0,0,0.08);border-color:#c7d2fe;}}
.tier-card.selected{{border-color:var(--primary-500);box-shadow:0 0 0 3px rgba(99,102,241,0.15),0 8px 24px rgba(0,0,0,0.08);}}
.tier-badge{{display:inline-block;padding:3px 10px;border-radius:10px;font-size:11px;font-weight:700;text-transform:uppercase;letter-spacing:0.5px;margin-bottom:8px;}}
.tier-range{{font-size:13px;color:var(--neutral-500);margin-bottom:4px;}}
.tier-num{{font-size:26px;font-weight:800;color:var(--neutral-900);}}
.tier-pct{{font-size:14px;font-weight:400;color:var(--neutral-400);}}
.tier-avg{{font-size:12px;color:var(--neutral-500);margin-top:6px;}}

/* Grid & Cards */
.grid-2{{display:grid;grid-template-columns:1fr 1fr;gap:20px;}}
.card{{background:#fff;border-radius:var(--radius-xl);padding:20px 22px;box-shadow:0 2px 8px rgba(0,0,0,0.04);border:1px solid var(--neutral-200);}}
.card-title{{font-size:12px;font-weight:700;color:var(--neutral-600);text-transform:uppercase;letter-spacing:0.5px;margin-bottom:14px;}}

/* Horizontal bars */
.hbar-row{{display:grid;grid-template-columns:80px 1fr auto;align-items:center;gap:10px;margin-bottom:10px;padding:4px 8px;border-radius:var(--radius-md);cursor:pointer;transition:all .2s ease;}}
.hbar-row:hover{{background:var(--primary-50);}}
.hbar-row.selected{{background:var(--primary-100);}}
.hbar-label{{font-size:12px;color:var(--neutral-700);text-align:right;white-space:nowrap;font-weight:500;}}
.hbar-track{{background:var(--neutral-100);border-radius:4px;height:24px;display:flex;align-items:center;}}
.hbar-fill{{height:100%;border-radius:4px;min-width:3px;transition:width .4s ease;}}
.hbar-val{{font-size:12px;color:var(--neutral-600);font-weight:600;white-space:nowrap;}}
.hbar-cnt{{font-size:11px;color:var(--neutral-400);font-weight:400;}}

/* Volume */
.hbar-row.hbar-vol{{grid-template-columns:95px 1fr;}}
.hbar-mono{{font-family:'SF Mono',Monaco,'Cascadia Code',Consolas,monospace;font-size:11px;color:var(--neutral-600);}}
.hbar-inline{{font-size:11px;color:var(--neutral-500);margin-left:8px;white-space:nowrap;font-weight:500;}}
.vol-xaxis{{display:flex;justify-content:space-between;margin-top:8px;padding-left:105px;font-size:10px;color:var(--neutral-400);}}
.vol-filter-label{{font-size:12px;font-weight:600;color:var(--primary-600);margin-left:8px;}}

/* Table */
.table-wrap{{border-radius:var(--radius-xl);overflow:hidden;box-shadow:0 2px 8px rgba(0,0,0,0.04);border:1px solid var(--neutral-200);max-height:2200px;overflow-y:auto;}}
table{{width:100%;border-collapse:collapse;font-size:13px;}}
thead th{{padding:12px 16px;font-size:11px;font-weight:700;color:var(--neutral-600);text-transform:uppercase;letter-spacing:0.5px;background:linear-gradient(to bottom,#f8fafc,#f1f5f9);border-bottom:2px solid var(--neutral-200);position:sticky;top:0;z-index:5;white-space:nowrap;}}
tbody td{{padding:12px 16px;color:var(--neutral-700);border-bottom:1px solid #f5f5f5;}}
tbody tr:hover td{{background:var(--primary-50);}}
.th-sort,.th-filter{{color:var(--neutral-400);cursor:pointer;font-size:11px;margin-left:2px;transition:color .15s;}}
.th-sort:hover,.th-filter:hover{{color:var(--primary-500);}}
.th-sort.active{{color:var(--primary-600);}}

.asin-link{{display:inline-flex;align-items:center;gap:4px;padding:4px 10px;background:linear-gradient(135deg,#eff6ff,#dbeafe);border:1px solid #bfdbfe;border-radius:6px;font-family:'SF Mono',Monaco,Consolas,monospace;font-size:11px;color:#1d4ed8;font-weight:600;text-decoration:none;transition:all .2s cubic-bezier(0.34,1.56,0.64,1);}}
.asin-link:hover{{background:linear-gradient(135deg,#dbeafe,#bfdbfe);border-color:#3b82f6;transform:translateY(-1px);box-shadow:0 4px 8px rgba(59,130,246,0.2);text-decoration:none;}}
.cell-title{{color:var(--neutral-500);font-size:12px;}}
.price-tag{{font-family:'SF Mono',Monaco,monospace;font-weight:700;font-size:12px;color:var(--neutral-800);}}
.rank-badge{{display:inline-flex;align-items:center;justify-content:center;min-width:28px;padding:2px 8px;border-radius:6px;font-weight:700;font-size:12px;background:var(--neutral-100);color:var(--neutral-600);}}
.st-ok{{color:var(--success);font-weight:600;}}
.st-unavailable{{color:var(--warning);}}
.st-not_found,.st-error{{color:var(--danger);}}

/* Filter popup */
.filter-popup{{display:none;position:absolute;background:#fff;border:1px solid var(--neutral-200);border-radius:var(--radius-lg);padding:16px;z-index:1000;min-width:220px;max-height:360px;overflow-y:auto;box-shadow:0 12px 24px rgba(0,0,0,0.15);}}
.filter-popup.show{{display:block;}}
.fp-hint{{font-size:11px;color:var(--neutral-400);margin-bottom:8px;}}
.filter-popup label{{display:block;padding:4px 0;font-size:12px;color:var(--neutral-700);cursor:pointer;}}
.filter-popup input[type="checkbox"]{{margin-right:6px;accent-color:var(--primary-500);}}
.filter-popup input[type="number"]{{width:80px;border:1px solid var(--neutral-200);padding:5px 8px;border-radius:var(--radius-md);font-size:12px;}}
.fp-actions{{margin-top:12px;display:flex;gap:8px;}}
.fp-btn{{padding:6px 16px;border-radius:var(--radius-md);border:none;cursor:pointer;font-size:12px;font-weight:600;transition:all .2s ease;}}
.fp-apply{{background:linear-gradient(135deg,#6366f1,#4f46e5);color:#fff;box-shadow:0 4px 12px rgba(99,102,241,0.3);}}
.fp-apply:hover{{background:linear-gradient(135deg,#4f46e5,#4338ca);}}
.fp-clear{{background:var(--neutral-100);color:var(--neutral-600);}}
.fp-clear:hover{{background:var(--neutral-200);}}
.fp-selectall{{font-weight:700;padding-bottom:4px;border-bottom:1px solid var(--neutral-100);margin-bottom:4px;}}
/* Insight box */
.insight-box{{background:#f8fafc;border:1px solid var(--neutral-200);border-radius:var(--radius-lg);padding:18px 22px;margin-top:18px;}}
.insight-title{{font-size:13px;font-weight:700;color:var(--neutral-800);text-transform:uppercase;letter-spacing:0.5px;margin-bottom:10px;padding-bottom:8px;border-bottom:1px solid var(--neutral-200);}}
.insight-box p{{font-size:12.5px;color:var(--neutral-700);line-height:1.7;margin-bottom:10px;}}
.insight-box p.cn{{font-style:italic;color:var(--neutral-500);font-size:12px;margin-bottom:8px;}}
.insight-box b{{color:var(--neutral-900);}}
.insight-note{{font-size:11px;color:var(--neutral-400);margin-top:10px;font-style:italic;padding-top:8px;border-top:1px solid var(--neutral-200);}}

/* Tier-brand collapsible */
.tier-brand-section{{margin-bottom:12px;}}
.tier-brand-header{{padding:10px 14px;font-size:13px;font-weight:600;color:var(--neutral-800);margin-bottom:4px;}}
.tier-brand-detail{{margin-left:20px;}}
.tier-brand-detail summary{{padding:6px 10px;font-size:12px;cursor:pointer;border-radius:6px;transition:background .15s;}}
.tier-brand-detail summary:hover{{background:var(--neutral-100);}}
.tb-brand{{font-weight:600;color:var(--neutral-800);}}
.tb-count{{font-size:11px;color:var(--neutral-400);margin-left:4px;}}
.tb-asins{{padding:6px 10px 10px 10px;font-family:'SF Mono',Monaco,Consolas,monospace;font-size:11px;color:var(--neutral-500);line-height:1.6;}}

canvas{{width:100%!important;}}

/* Action buttons */
.table-actions{{display:flex;gap:8px;margin-bottom:12px;}}
.action-btn{{display:inline-flex;align-items:center;gap:6px;padding:8px 16px;background:var(--primary-50);border:1px solid #c7d2fe;border-radius:var(--radius-md);font-size:12px;font-weight:600;color:var(--primary-700);cursor:pointer;transition:all .2s cubic-bezier(0.34,1.56,0.64,1);}}
.action-btn:hover{{background:var(--primary-100);transform:translateY(-1px);box-shadow:0 4px 12px rgba(99,102,241,0.2);}}
.action-btn.copied{{background:#dcfce7;border-color:#86efac;color:#166534;}}
.toast{{position:fixed;top:20px;left:50%;transform:translateX(-50%) translateY(-100px);background:#1E293B;color:#fff;padding:10px 24px;border-radius:10px;font-size:13px;font-weight:600;z-index:9999;opacity:0;transition:all .3s ease;pointer-events:none;}}
.toast.show{{transform:translateX(-50%) translateY(0);opacity:1;}}

/* Tier SVG chart */
.tier-chart-card{{padding:20px 24px;}}
.tier-svg-wrap{{display:flex;align-items:flex-end;justify-content:center;gap:40px;padding:20px 0 0;position:relative;}}
.tier-bar-col{{display:flex;flex-direction:column;align-items:center;cursor:pointer;transition:transform .2s cubic-bezier(0.34,1.56,0.64,1);}}
.tier-bar-col:hover{{transform:scale(1.03);}}
.tier-bar-col.selected .tier-bar-rect{{outline:3px solid var(--primary-500);outline-offset:3px;}}
.tier-bar-count{{font-size:18px;font-weight:800;color:var(--neutral-900);margin-bottom:8px;}}
.tier-bar-rect{{border-radius:8px 8px 0 0;min-width:100px;transition:height .5s ease;}}
.tier-bar-label{{margin-top:12px;font-size:12px;font-weight:600;color:var(--neutral-600);}}

/* Bubble cloud */
.bubble-cloud{{display:flex;flex-wrap:wrap;gap:10px;padding:8px 0;}}
.bubble{{display:inline-flex;align-items:center;gap:6px;padding:8px 16px;border-radius:9999px;color:#fff;font-size:12px;font-weight:600;cursor:pointer;transition:all .2s cubic-bezier(0.34,1.56,0.64,1);box-shadow:0 2px 8px rgba(0,0,0,0.12);}}
.bubble:hover{{transform:scale(1.08);box-shadow:0 4px 16px rgba(0,0,0,0.2);}}
.bubble.selected{{outline:3px solid var(--neutral-900);outline-offset:2px;transform:scale(1.08);}}
.bubble b{{font-weight:800;opacity:.9;background:rgba(255,255,255,0.25);padding:1px 6px;border-radius:8px;font-size:11px;}}

@media(max-width:900px){{.grid-2{{grid-template-columns:1fr;}} .tier-svg-wrap{{gap:20px;}}}}
</style>
</head>
<body>
<div class="container">
  <div class="hero">
    <h1>{title}</h1>
    {'<div class="sub">' + subtitle + '</div>' if subtitle else ''}
  </div>
  <div class="tabs">{tabs_html}</div>
  {panels_html}
</div>
<div class="filter-popup" id="filterPopup"></div>
<div class="toast" id="toast"></div>
<script>
const rowData={row_json};
const volumeSeries={vol_j};
const TIER_BOUNDS={tier_j};
const COLORS=['#6366f1','#8b5cf6','#ec4899','#f59e0b','#10b981','#06b6d4','#f97316','#64748b','#3b82f6','#22c55e'];
const TIER_COLORS=['#10b981','#3b82f6','#ec4899'];

function computeTierBounds(mk){{return TIER_BOUNDS[mk]||[];}}
function getPriceTier(price,mk){{
  if(!price)return null;
  const t=computeTierBounds(mk);
  for(let i=0;i<t.length;i++){{
    if(i===t.length-1){{if(price>=t[i].lo)return t[i].name;}}
    else{{if(price>=t[i].lo&&price<t[i].hi)return t[i].name;}}
  }}
  return t.length?t[t.length-1].name:null;
}}

// Tabs
document.querySelectorAll('.tab').forEach(t=>t.addEventListener('click',()=>{{
  document.querySelectorAll('.tab').forEach(x=>x.classList.remove('active'));
  document.querySelectorAll('.tab-content').forEach(x=>{{x.classList.remove('active');x.style.display='none';}});
  t.classList.add('active');
  const p=document.querySelector('.tab-content[data-market="'+t.dataset.market+'"]');
  if(p){{p.classList.add('active');p.style.display='block';}}
}}));

// Volume chart
function renderVolChart(mk,fType,fVal){{
  const c=document.getElementById('vol-chart-'+mk),lb=document.getElementById('vol-filter-'+mk);
  if(!c)return;
  let items=rowData.filter(r=>r.market===mk&&r.search_volume);
  if(fType==='price'){{items=items.filter(r=>getPriceTier(r.price,mk)===fVal);if(lb)lb.textContent='[ '+fVal+' ]';}}
  else if(fType==='brand'){{items=items.filter(r=>r.brand===fVal);if(lb)lb.textContent='[ '+fVal+' ]';}}
  else{{if(lb)lb.textContent='';}}
  items.sort((a,b)=>(b.search_volume||0)-(a.search_volume||0));
  items=items.slice(0,15);
  if(!items.length){{c.innerHTML='<div style="color:#a3a3a3;padding:24px;font-size:12px">No data</div>';return;}}
  const mx=items[0].search_volume||1;
  let h='';
  items.forEach((r,i)=>{{
    const pct=(r.search_volume/mx*100).toFixed(1);
    h+='<div class="hbar-row hbar-vol"><div class="hbar-label hbar-mono">'+r.asin+'</div><div class="hbar-track"><div class="hbar-fill" style="width:'+pct+'%;background:'+COLORS[i%10]+'"></div><span class="hbar-inline">'+( r.brand||'')+'</span></div></div>';
  }});

  c.innerHTML=h;
}}
document.querySelectorAll('[id^="vol-chart-"]').forEach(el=>renderVolChart(el.id.replace('vol-chart-',''),null,null));

// Tier bar chart (HTML-based, crisp)
document.querySelectorAll('[id^="tier-svg-"]').forEach(wrap=>{{
  const mk=wrap.id.replace('tier-svg-','');
  const t=computeTierBounds(mk);
  const prices=rowData.filter(r=>r.market===mk&&r.price).map(r=>r.price);
  const counts=t.map((tier,i)=>prices.filter(p=>p>=tier.lo&&(i===2||p<tier.hi)).length);
  const maxC=Math.max(...counts)||1;
  let html='';
  const maxH=160;
  counts.forEach((c,i)=>{{
    const barH=Math.max(Math.round(c/maxC*maxH),6);
    const lbl=t[i].name==='Entry'?'$0\u2013$'+Math.round(t[i].hi):t[i].name==='Mid-tier'?'$'+Math.round(t[i].lo)+'\u2013$'+Math.round(t[i].hi):'$'+Math.round(t[i].lo)+'+';
    html+='<div class="tier-bar-col clickable-bar" data-filter-type="price" data-filter-value="'+t[i].name+'" data-market="'+mk+'">';
    html+='<div class="tier-bar-count">'+c+'</div>';
    html+='<div class="tier-bar-rect" style="height:'+barH+'px;width:100px;background:linear-gradient(180deg,'+TIER_COLORS[i]+','+TIER_COLORS[i]+'88)"></div>';
    html+='<div class="tier-bar-label">'+lbl+'</div>';
    html+='</div>';
  }});
  wrap.innerHTML=html;
}});

// Brand bubble cloud
document.querySelectorAll('[id^="bubble-"]').forEach(container=>{{
  const mk=container.id.replace('bubble-','');
  const brandMap={{}};
  rowData.filter(r=>r.market===mk&&r.brand).forEach(r=>{{brandMap[r.brand]=(brandMap[r.brand]||0)+1;}});
  const sorted=Object.entries(brandMap).sort((a,b)=>b[1]-a[1]).slice(0,20);
  let html='';
  sorted.forEach(([brand,cnt],i)=>{{
    const color=COLORS[i%10];
    html+='<span class="bubble clickable-bar" data-filter-type="brand" data-filter-value="'+brand.replace(/"/g,'&quot;')+'" data-market="'+mk+'" style="background:'+color+'">'+brand+' <b>'+cnt+'</b></span>';
  }});
  container.innerHTML=html;
}});

// Donut charts
function drawDonut(canvasId, dataMap, title){{
  const canvas=document.getElementById(canvasId);
  if(!canvas)return;
  const entries=Object.entries(dataMap).filter(e=>e[1]>0).sort((a,b)=>b[1]-a[1]);
  if(!entries.length)return;
  const total=entries.reduce((s,e)=>s+e[1],0);
  // Top 10 + Others
  let items=entries.slice(0,10);
  const othersVal=entries.slice(10).reduce((s,e)=>s+e[1],0);
  if(othersVal>0)items.push(['OTHERS',othersVal]);

  const ctx=canvas.getContext('2d');
  const W=canvas.offsetWidth||350,H=260;
  canvas.width=W;canvas.height=H;
  const cx=W*0.38,cy=H/2,R=Math.min(cx-10,cy-10),r=R*0.55;
  let angle=-Math.PI/2;

  items.forEach((item,i)=>{{
    const slice=item[1]/total*Math.PI*2;
    const color=i<10?COLORS[i%10]:'#d4d4d4';
    ctx.beginPath();ctx.moveTo(cx,cy);
    ctx.arc(cx,cy,R,angle,angle+slice);ctx.closePath();
    ctx.fillStyle=color;ctx.fill();
    angle+=slice;
  }});
  // Inner circle (donut hole)
  ctx.beginPath();ctx.arc(cx,cy,r,0,Math.PI*2);
  ctx.fillStyle='#ffffff';ctx.fill();

  // Legend on right
  const legX=cx+R+20,legY=20;
  ctx.font='11px Inter,sans-serif';
  items.forEach((item,i)=>{{
    const y=legY+i*20;
    const pct=(item[1]/total*100).toFixed(1);
    const color=i<10?COLORS[i%10]:'#d4d4d4';
    ctx.fillStyle=color;
    ctx.beginPath();ctx.arc(legX,y+5,5,0,Math.PI*2);ctx.fill();
    ctx.fillStyle='#404040';ctx.font='600 11px Inter';
    ctx.fillText(item[0]+' ('+pct+'%)',legX+12,y+9);
  }});
}}

// Render donut charts per market
document.querySelectorAll('[id^="donut-sv-"]').forEach(c=>{{
  const mk=c.id.replace('donut-sv-','');
  const brandSV={{}};
  rowData.filter(r=>r.market===mk&&r.search_volume&&r.brand).forEach(r=>{{
    brandSV[r.brand]=(brandSV[r.brand]||0)+r.search_volume;
  }});
  drawDonut(c.id,brandSV);
}});
document.querySelectorAll('[id^="donut-cnt-"]').forEach(c=>{{
  const mk=c.id.replace('donut-cnt-','');
  const brandCnt={{}};
  rowData.filter(r=>r.market===mk&&r.price&&r.brand).forEach(r=>{{
    brandCnt[r.brand]=(brandCnt[r.brand]||0)+1;
  }});
  drawDonut(c.id,brandCnt);
}});

// Click interaction
document.querySelectorAll('.clickable-bar').forEach(bar=>bar.addEventListener('click',()=>{{
  const mk=bar.dataset.market,type=bar.dataset.filterType,val=bar.dataset.filterValue;
  const was=bar.classList.contains('selected');
  document.querySelectorAll('.clickable-bar[data-market="'+mk+'"]').forEach(b=>b.classList.remove('selected'));
  if(!was){{bar.classList.add('selected');renderVolChart(mk,type,val);}}
  else renderVolChart(mk,null,null);
}}));

// Table sort
document.querySelectorAll('.th-sort').forEach(el=>el.addEventListener('click',e=>{{
  e.stopPropagation();
  const th=el.closest('th'),table=th.closest('table'),tbody=table.querySelector('tbody');
  const col=parseInt(el.dataset.col),rows=Array.from(tbody.querySelectorAll('tr'));
  const asc=!el.classList.contains('asc');
  table.querySelectorAll('.th-sort').forEach(s=>{{s.classList.remove('asc','desc','active');s.innerHTML='&#8597;';}});
  el.classList.add(asc?'asc':'desc','active');el.innerHTML=asc?'&#9650;':'&#9660;';
  rows.sort((a,b)=>{{let va=a.cells[col]?.textContent.trim()||'',vb=b.cells[col]?.textContent.trim()||'';const na=parseFloat(va.replace(/[^\\d.\\-]/g,'')),nb=parseFloat(vb.replace(/[^\\d.\\-]/g,''));if(!isNaN(na)&&!isNaN(nb))return asc?na-nb:nb-na;return asc?va.localeCompare(vb):vb.localeCompare(va);}});
  rows.forEach(r=>tbody.appendChild(r));
}}));

// Table filter
const popup=document.getElementById('filterPopup');let fCol=null,fTable=null;
document.querySelectorAll('.th-filter').forEach(el=>el.addEventListener('click',e=>{{
  e.stopPropagation();
  const th=el.closest('th'),table=th.closest('table'),col=parseInt(el.dataset.col);
  fCol=col;fTable=table;
  const rows=Array.from(table.querySelector('tbody').querySelectorAll('tr'));
  const vals=new Set();let isNum=true;
  rows.forEach(r=>{{const v=r.cells[col]?.textContent.trim()||'';vals.add(v);if(v&&v!=='\\u2014'&&isNaN(parseFloat(v.replace(/[^\\d.\\-]/g,''))))isNum=false;}});
  let h='';
  if(isNum&&vals.size>5){{
    h='<div class="fp-hint">Numeric range</div><label>Min: <input type="number" id="fMin" step="any"></label><label style="margin-top:6px">Max: <input type="number" id="fMax" step="any"></label><div class="fp-actions"><button class="fp-btn fp-apply" onclick="applyNum()">Apply</button><button class="fp-btn fp-clear" onclick="clearF()">Clear</button></div>';
  }}else{{
    const sorted=Array.from(vals).sort();
    h='<div class="fp-hint">Select values</div><label class="fp-selectall"><input type="checkbox" id="fpSelAll" checked onchange="toggleAll(this)"> <strong>Select All</strong></label>';
    sorted.forEach(v=>{{h+='<label><input type="checkbox" class="fp-cb" value="'+(v.replace(/"/g,'&quot;'))+'" checked> '+(v||'(empty)')+'</label>';}});
    h+='<div class="fp-actions"><button class="fp-btn fp-apply" onclick="applyTxt()">Apply</button><button class="fp-btn fp-clear" onclick="clearF()">Clear</button></div>';
  }}
  popup.innerHTML=h;
  const rect=el.getBoundingClientRect();
  popup.style.top=(rect.bottom+window.scrollY+4)+'px';
  popup.style.left=Math.min(rect.left+window.scrollX,window.innerWidth-250)+'px';
  popup.classList.add('show');
}}));
document.addEventListener('click',e=>{{if(!popup.contains(e.target)&&!e.target.classList.contains('th-filter'))popup.classList.remove('show');}});
function toggleAll(el){{popup.querySelectorAll('.fp-cb').forEach(cb=>cb.checked=el.checked);}}
function applyNum(){{const mn=parseFloat(document.getElementById('fMin').value),mx=parseFloat(document.getElementById('fMax').value);Array.from(fTable.querySelector('tbody').rows).forEach(r=>{{const v=parseFloat((r.cells[fCol]?.textContent||'').replace(/[^\\d.\\-]/g,''));let ok=true;if(!isNaN(mn)&&(isNaN(v)||v<mn))ok=false;if(!isNaN(mx)&&(isNaN(v)||v>mx))ok=false;r.style.display=ok?'':'none';}});popup.classList.remove('show');}}
function applyTxt(){{const ck=new Set();popup.querySelectorAll('.fp-cb:checked').forEach(c=>ck.add(c.value));Array.from(fTable.querySelector('tbody').rows).forEach(r=>{{r.style.display=ck.has(r.cells[fCol]?.textContent.trim()||'')?'':'none';}});popup.classList.remove('show');}}
function clearF(){{Array.from(fTable.querySelector('tbody').rows).forEach(r=>r.style.display='');popup.classList.remove('show');}}

// === Copy Volume ASINs ===
function copyVolAsins(mk){{
  const container=document.getElementById('vol-chart-'+mk);
  if(!container)return;
  const asins=Array.from(container.querySelectorAll('.hbar-mono')).map(el=>el.textContent.trim()).filter(a=>a.length===10);
  if(!asins.length){{showToast('No ASINs to copy');return;}}
  navigator.clipboard.writeText(asins.join(',')).then(()=>showToast('Copied '+asins.length+' ASINs')).catch(()=>showToast('Copy failed'));
}}

// === Copy Table ASINs ===
function showToast(msg){{
  const t=document.getElementById('toast');
  t.textContent=msg;t.classList.add('show');
  setTimeout(()=>t.classList.remove('show'),2000);
}}
function copyAsins(mk){{
  const table=document.getElementById('table-'+mk);
  if(!table)return;
  const rows=Array.from(table.querySelector('tbody').querySelectorAll('tr'));
  const asins=[];
  rows.forEach(r=>{{
    if(r.style.display==='none')return;
    const link=r.cells[0]?.querySelector('a');
    const asin=link?link.textContent.trim():(r.cells[0]?.textContent.trim()||'');
    if(asin&&asin.length===10)asins.push(asin);
  }});
  if(!asins.length){{showToast('No ASINs to copy');return;}}
  navigator.clipboard.writeText(asins.join(',')).then(()=>{{
    showToast('Copied '+asins.length+' ASINs');
  }}).catch(()=>showToast('Copy failed'));
}}

// === Export Excel ===
function exportExcel(mk){{
  const table=document.getElementById('table-'+mk);
  if(!table)return;
  const headers=Array.from(table.querySelectorAll('thead th')).map(th=>th.textContent.replace(/[\\u25B2\\u25BC\\u2195\\u25BC]/g,'').trim());
  const rows=Array.from(table.querySelector('tbody').querySelectorAll('tr'));
  let csv='\\uFEFF'+headers.join('\\t')+'\\n';
  rows.forEach(r=>{{
    if(r.style.display==='none')return;
    const cells=Array.from(r.cells).map(td=>td.textContent.trim().replace(/\\t/g,' '));
    csv+=cells.join('\\t')+'\\n';
  }});
  const blob=new Blob([csv],{{type:'application/vnd.ms-excel;charset=utf-8'}});
  const url=URL.createObjectURL(blob);
  const a=document.createElement('a');
  a.href=url;a.download='ASIN_Detail_'+mk+'.xls';
  document.body.appendChild(a);a.click();document.body.removeChild(a);
  URL.revokeObjectURL(url);
  showToast('Exported '+mk+' table');
}}
</script>
</body>
</html>'''
